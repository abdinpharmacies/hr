import base64
import binascii
import math
import re
from io import BytesIO
from zipfile import BadZipFile, ZipFile

from openpyxl import Workbook, load_workbook

from odoo import fields, models, _
from odoo.exceptions import UserError


class AbTransferSmartStockImport(models.TransientModel):
    _name = "ab_transfer_smart_stock_import"
    _description = "Import Destination Stock"

    store_id = fields.Many2one("ab_store", string="Destination Branch", required=True)
    file_data = fields.Binary(string="Excel File", attachment=False)
    filename = fields.Char(string="Filename")
    confirm_replace = fields.Boolean(string="Replace the entire destination stock snapshot")
    validation_summary = fields.Text(string="Validation Summary", readonly=True)
    template_data = fields.Binary(attachment=False, readonly=True)
    template_filename = fields.Char(readonly=True)

    def _check_import_access(self):
        self.ensure_one()
        self.check_access("write")
        self.env["ab_transfer_smart_cache_state"]._check_branch_access(self.store_id, importing=True)
        if not self.store_id or not self.store_id.exists():
            raise UserError(_("Select one destination branch."))
        if not self.store_id.eplus_serial:
            raise UserError(_("The destination branch needs a valid E-Plus ID."))

    def action_download_template(self):
        self._check_import_access()
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Destination Stock"
        sheet.append(["Product Code", "On-hand Quantity", "Pending Quantity"])
        sheet.column_dimensions["A"].width = 24
        sheet.column_dimensions["B"].width = 24
        sheet.column_dimensions["C"].width = 24
        for row in sheet.iter_rows(min_row=2, max_row=1001, max_col=1):
            row[0].number_format = "@"
        stream = BytesIO()
        workbook.save(stream)
        self.write({"template_data": base64.b64encode(stream.getvalue()),
                    "template_filename": "destination_stock_template.xlsx"})
        return {"type": "ir.actions.act_url", "target": "self",
                "url": "/web/content/ab_transfer_smart_stock_import/%s/template_data/destination_stock_template.xlsx?download=true" % self.id}

    def _parse_file(self):
        self._check_import_access()
        if not self.file_data:
            raise UserError(_("Upload an XLSX workbook first."))
        try:
            raw = base64.b64decode(self.file_data, validate=True)
            if len(raw) > 10 * 1024 * 1024:
                raise ValueError()
            with ZipFile(BytesIO(raw)) as archive:
                if sum(info.file_size for info in archive.infolist()) > 100 * 1024 * 1024:
                    raise ValueError()
            workbook = load_workbook(BytesIO(raw), read_only=True, data_only=False)
        except (ValueError, binascii.Error, BadZipFile, KeyError, OSError):
            raise UserError(_("Upload a valid XLSX workbook smaller than 10 MB.")) from None
        try:
            if len(workbook.worksheets) != 1:
                raise UserError(_("The workbook must contain exactly one worksheet."))
            sheet = workbook.worksheets[0]
            if sheet.max_row > 100001 or sheet.max_column > 20:
                raise UserError(_("The workbook exceeds the supported row or column limit."))
            rows = iter(sheet.iter_rows())
            header = next(rows, ())
            columns = [str(cell.value or "").strip() for cell in header]
            required = ["Product Code", "On-hand Quantity"]
            if any(columns.count(name) != 1 for name in required) or columns.count("Pending Quantity") > 1:
                raise UserError(_("Required columns: Product Code, On-hand Quantity; optional: Pending Quantity."))
            code_index, qty_index = [columns.index(name) for name in required]
            pending_index = columns.index("Pending Quantity") if "Pending Quantity" in columns else None
            parsed = {}
            for row_number, row in enumerate(rows, 2):
                if all(cell.value is None for cell in row):
                    continue
                if any(cell.data_type in ("f", "e") for cell in row):
                    raise UserError(_("Row %s contains a formula or Excel error.") % row_number)
                cell = row[code_index]
                value = cell.value
                if isinstance(value, str):
                    code = value
                elif isinstance(value, (int, float)) and not isinstance(value, bool) and math.isfinite(value) and value == int(value):
                    code = str(int(value))
                    if re.fullmatch(r"0+", cell.number_format or ""):
                        code = code.zfill(len(cell.number_format))
                else:
                    code = ""
                if not code or code in parsed:
                    raise UserError(_("Row %s has an empty or duplicate product code.") % row_number)
                qty = self._quantity(row[qty_index].value, row_number)
                pending = self._quantity(row[pending_index].value, row_number, pending=True) if pending_index is not None else 0.0
                if not math.isfinite(qty + pending):
                    raise UserError(_("Row %s has an invalid quantity.") % row_number)
                parsed[code] = (qty, pending)
            if not parsed:
                raise UserError(_("The workbook contains no product rows."))
        finally:
            workbook.close()
        products = self.env["ab_product"].with_context(active_test=False).search([("code", "in", list(parsed))])
        by_code = {}
        for product in products:
            by_code.setdefault(product.code, []).append(product)
        result = []
        serials = set()
        for code, (qty, pending) in parsed.items():
            matches = by_code.get(code, [])
            if len(matches) != 1:
                raise UserError(_("Product code %s is missing or ambiguous.") % code)
            try:
                serial = int(matches[0].eplus_serial)
                if serial <= 0 or str(serial) != str(matches[0].eplus_serial).strip():
                    raise ValueError()
            except (TypeError, ValueError):
                raise UserError(_("Product code %s needs a valid E-Plus ID.") % code) from None
            if serial in serials:
                raise UserError(_("Multiple rows resolve to E-Plus ID %s.") % serial)
            serials.add(serial)
            result.append({"store_id": self.store_id.id, "product_eplus_serial": serial,
                           "stock_qty": qty + pending, "stock_pending_qty": pending,
                           "cache_date": self.env["ab_transfer_smart_cache_state"]._business_date()})
        all_products = self.env["ab_product"].with_context(active_test=False).search([("eplus_serial", "in", list(serials))])
        if len(all_products) != len(serials):
            raise UserError(_("The imported E-Plus IDs are ambiguous in the product master."))
        return result

    def _quantity(self, value, row_number, pending=False):
        if pending and value in (None, ""):
            return 0.0
        try:
            if isinstance(value, bool) or value in (None, ""):
                raise ValueError()
            quantity = float(value)
            if not math.isfinite(quantity) or (pending and quantity < 0):
                raise ValueError()
            return quantity
        except (TypeError, ValueError, OverflowError):
            raise UserError(_("Row %s has an invalid quantity.") % row_number) from None

    def action_refresh_sales(self):
        self._check_import_access()
        self.env["ab_transfer_smart_sales_cache"].refresh_store_cache(self.store_id, force=True)
        self.validation_summary = _("Destination sales refreshed successfully for the current Cairo business date.")
        return self._reopen()

    def action_validate_file(self):
        rows = self._parse_file()
        self.validation_summary = _("Validated %s products. Omitted products will have zero stock. Pending is included once.") % len(rows)
        return self._reopen()

    def action_import(self):
        rows = self._parse_file()
        if not self.confirm_replace:
            raise UserError(_("Confirm replacement of the entire destination stock snapshot."))
        State = self.env["ab_transfer_smart_cache_state"]
        State._lock_branches(self.store_id)
        self._check_import_access()
        self.env["ab_transfer_smart_stock_cache"]._replace_store_stock_snapshot(
            self.store_id, rows, origin="excel", filename=self.filename, actor_id=self.env.uid,
        )
        ready = self.env["ab_transfer_smart_sales_cache"].read_store_sales_context(self.store_id, [])["complete"]
        self.validation_summary = (
            _("Stock imported. Valid until the next Cairo midnight. Destination sales are ready.") if ready else
            _("Stock imported. Calculation is blocked until this branch has a verified current-day destination sales refresh.")
        )
        self.confirm_replace = False
        return self._reopen()

    def _reopen(self):
        return {"type": "ir.actions.act_window", "res_model": self._name, "res_id": self.id,
                "view_mode": "form", "target": "new"}
