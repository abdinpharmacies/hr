import base64
import binascii
from io import BytesIO
from zipfile import BadZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from odoo import _, api, fields, models
from odoo.exceptions import AccessError, UserError

from .product_seo import SEO_LANGUAGES


HEADERS = {
    "كود المنتج": "product_code",
    "عنوان Meta": "meta_title",
    "وصف Meta": "meta_description",
    "الكلمات المفتاحية": "keyword_text",
    "الرابط المختصر": "seo_name",
    "وصف مختصر": "short_description",
    "الوصف العام": "public_description",
    "المادة الفعالة": "active_ingredient",
    "التحذيرات": "warnings",
    "موانع الاستخدام": "contraindications",
    "التخزين": "storage",
}
MAX_FILE_BYTES = 10 * 1024 * 1024
MAX_DATA_ROWS = 5000


class AbProductSeoXlsxImportWizard(models.TransientModel):
    _name = "ab_product_seo_xlsx_import_wizard"
    _description = "Import Product SEO from Excel"

    file = fields.Binary(required=True, attachment=False, string="Excel File")
    filename = fields.Char(string="Filename")
    template_file = fields.Binary(default=lambda self: self._get_template_file(), attachment=False, string="Excel Template")
    template_filename = fields.Char(default="product_seo_import_template.xlsx")
    lang_code = fields.Selection(SEO_LANGUAGES, required=True, default="en_US", string="Content Language")
    website_id = fields.Many2one("website", string="Website")
    import_outcome = fields.Selection(
        [("success", "Imported"), ("failed", "Not Imported")], readonly=True,
    )
    import_summary = fields.Text(string="Import Summary", readonly=True)
    import_error = fields.Text(string="Import Error", readonly=True)
    report_line_ids = fields.One2many(
        "ab_product_seo_xlsx_import_line", "wizard_id", string="Import Report", readonly=True,
    )

    @api.onchange("file", "lang_code", "website_id")
    def _onchange_import_configuration(self):
        self.import_outcome = False
        self.import_summary = False
        self.import_error = False
        self.report_line_ids = [fields.Command.clear()]

    def action_import_report(self):
        self.ensure_one()
        self.check_access("write")
        if not self.env.user.has_group("ab_website_seo_optimization.group_ab_website_seo_optimization_manager"):
            raise AccessError(_("Only SEO Managers can import product SEO."))
        rows = []
        try:
            with self.env.cr.savepoint():
                rows = self._read_rows()
                action = self.action_import()
        except AccessError:
            raise
        except UserError as error:
            lines = [{
                "row_number": str(number),
                "product_code": code,
                "status": "not_imported",
                "reason": _("No rows were saved because the file contains an error."),
            } for number, code, _values in rows]
            if not lines:
                lines = [{"status": "error", "reason": str(error)}]
            return self._show_import_report(
                "failed",
                _("Imported: 0. Not imported: %(count)s.") % {"count": len(rows)}
                if rows else _("Import failed. No products were imported."),
                lines, error=str(error),
            )
        records = self.env["ab.product.seo"].search(action["domain"])
        by_code = {record.ab_product_id.code: record for record in records}
        lines = [{
            "row_number": str(number),
            "product_code": code,
            "product_name": by_code[code].product_template_id.display_name,
            "seo_id": by_code[code].id,
            "status": "imported",
            "reason": _("SEO draft saved for review."),
        } for number, code, _values in rows]
        return self._show_import_report(
            "success", _("Imported: %(count)s. Not imported: 0.") % {"count": len(lines)}, lines,
        )

    def _show_import_report(self, outcome, summary, lines, error=False):
        self.write({
            "import_outcome": outcome,
            "import_summary": summary,
            "import_error": error,
            "report_line_ids": [fields.Command.clear()] + [fields.Command.create(line) for line in lines],
        })
        return {
            "type": "ir.actions.act_window",
            "name": _("Import SEO from Excel"),
            "res_model": self._name,
            "res_id": self.id,
            "view_mode": "form",
            "view_id": self.env.ref("ab_website_seo_optimization.view_ab_product_seo_xlsx_import_wizard_form").id,
            "target": "new",
        }

    def action_open_imported_drafts(self):
        self.ensure_one()
        self.check_access("read")
        return {
            "type": "ir.actions.act_window",
            "name": _("Imported SEO Drafts"),
            "res_model": "ab.product.seo",
            "view_mode": "list,form",
            "domain": [("id", "in", self.report_line_ids.mapped("seo_id").ids)],
        }

    def _get_template_file(self):
        workbook = Workbook()
        workbook.active.append(list(HEADERS))
        workbook.active.column_dimensions["A"].number_format = "@"
        output = BytesIO()
        workbook.save(output)
        return base64.b64encode(output.getvalue())

    def _read_rows(self):
        self.ensure_one()
        if not (self.filename or "").lower().endswith(".xlsx"):
            raise UserError(_("Upload an .xlsx Excel file."))
        try:
            content = base64.b64decode(self.file, validate=True)
        except (binascii.Error, ValueError):
            raise UserError(_("The Excel file could not be decoded.")) from None
        if len(content) > MAX_FILE_BYTES:
            raise UserError(_("The Excel file must be 10 MB or smaller."))
        try:
            workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
            sheet = workbook.active
            rows = sheet.iter_rows(values_only=True)
            header = next(rows, ())
            header_names = [str(value).strip() if value is not None else "" for value in header]
            nonempty_headers = [name for name in header_names if name]
            if len(nonempty_headers) != len(set(nonempty_headers)):
                raise UserError(_("The Excel file has duplicate column headings."))
            missing = [name for name in HEADERS if name not in header_names]
            if missing:
                raise UserError(_("Missing Excel columns: %s") % ", ".join(missing))
            positions = {name: header_names.index(name) for name in HEADERS}
            parsed = []
            seen_codes = set()
            for row_number, values in enumerate(rows, start=2):
                if row_number > MAX_DATA_ROWS + 1:
                    raise UserError(_("The Excel file can contain at most %s product rows.") % MAX_DATA_ROWS)
                if not any(value is not None and str(value).strip() for value in values):
                    continue
                raw_code = values[positions["كود المنتج"]] if positions["كود المنتج"] < len(values) else None
                if type(raw_code) in (int, float) and raw_code >= 0 and raw_code < 10**15 and int(raw_code) == raw_code:
                    raw_code = str(int(raw_code))
                if not isinstance(raw_code, str) or not raw_code.strip():
                    raise UserError(_("Row %(row)s needs a text product code.") % {"row": row_number})
                code = raw_code.strip()
                if code in seen_codes:
                    raise UserError(_("Product code %(code)s appears more than once in the Excel file.") % {"code": code})
                seen_codes.add(code)
                content_values = {}
                for heading, field_name in HEADERS.items():
                    if field_name == "product_code":
                        continue
                    position = positions[heading]
                    value = values[position] if position < len(values) else None
                    if value is not None and str(value).strip():
                        if not isinstance(value, str):
                            raise UserError(_("Row %(row)s has a non-text value in %(column)s.") % {
                                "row": row_number, "column": heading,
                            })
                        content_values[field_name] = value.strip()
                if not content_values.get("meta_title") or not content_values.get("meta_description"):
                    raise UserError(_("Row %(row)s needs both Meta Title and Meta Description.") % {"row": row_number})
                parsed.append((row_number, code, content_values))
            if not parsed:
                raise UserError(_("The Excel file has no product rows."))
            return parsed
        except (BadZipFile, InvalidFileException, OSError, ValueError, KeyError, IndexError) as error:
            raise UserError(_("The Excel file is invalid: %s") % error) from error
        finally:
            if "workbook" in locals():
                workbook.close()

    def action_import(self):
        self.ensure_one()
        if not self.env.user.has_group("ab_website_seo_optimization.group_ab_website_seo_optimization_manager"):
            raise AccessError(_("Only SEO Managers can import product SEO."))
        if self.website_id and self.website_id.company_id and self.website_id.company_id not in self.env.companies:
            raise AccessError(_("The selected website is outside your allowed companies."))
        rows = self._read_rows()
        codes = [code for _number, code, _values in rows]
        products = self.env["ab_product"].search([("code", "in", codes)])
        products_by_code = {}
        for product in products:
            products_by_code.setdefault(product.code, []).append(product)
        templates = self.env["product.template"].search([
            ("ab_product_id", "in", products.ids),
            ("active", "=", True),
            ("sale_ok", "=", True),
            ("is_published", "=", True),
            ("company_id", "in", [False] + self.env.companies.ids),
        ])
        templates_by_product = {}
        for template in templates:
            if template.website_id and template.website_id != self.website_id:
                continue
            templates_by_product.setdefault(template.ab_product_id.id, []).append(template)
        seo_records = self.env["ab.product.seo"].search([
            ("ab_product_id", "in", products.ids),
            ("website_id", "=", self.website_id.id or False),
            ("active", "=", True),
        ])
        seo_by_product = {record.ab_product_id.id: record for record in seo_records}
        prepared = []
        for number, code, values in rows:
            matches = products_by_code.get(code, [])
            if len(matches) != 1:
                raise UserError(_("Row %(row)s: product code %(code)s must match exactly one product.") % {
                    "row": number, "code": code,
                })
            product = matches[0]
            available = templates_by_product.get(product.id, [])
            if len(available) != 1:
                raise UserError(_("Row %(row)s: product code %(code)s must have exactly one published website product.") % {
                    "row": number, "code": code,
                })
            template = available[0]
            record = seo_by_product.get(product.id)
            if self.website_id and template.company_id and template.company_id != self.website_id.company_id:
                raise AccessError(_("The selected website and product belong to different companies."))
            if record and (record.product_template_id != template or record.state in ("under_review", "approved")):
                raise UserError(_("Row %(row)s: SEO record for %(code)s needs manual review before import.") % {
                    "row": number, "code": code,
                })
            target_company = template.company_id or self.website_id.company_id or self.env.company
            if record and record.company_id != target_company:
                raise AccessError(_("The SEO record company does not match the product company."))
            prepared.append((product, template, record, target_company, values))
        imported = self.env["ab.product.seo"]
        for product, template, record, company, values in prepared:
            if not record:
                record = imported.create({
                    "ab_product_id": product.id,
                    "product_template_id": template.id,
                    "company_id": company.id,
                    "website_id": self.website_id.id or False,
                    "lang_mode": self.lang_code,
                })
            record._ensure_default_translations()
            translation = record.translation_ids.filtered(lambda item: item.lang_code == self.lang_code)
            translation.check_access("write")
            translation.sudo().write({
                "seo_template_id": False,
                "template_version": 0,
                "template_snapshot": False,
                "trusted_facts": False,
                "structured_sections": False,
                "validation_status": "pending",
                "validation_report": False,
                "generation_provider_id": False,
                "generation_model": False,
                "generated_at": False,
                "enrichment_item_id": False,
                "review_required": False,
                "reviewed_by": False,
                "reviewed_at": False,
                "prompt_version": False,
                "search_phrases": False,
                "structured_bullets": False,
            })
            if "public_description" in values:
                values = dict(values, public_description=record._html_paragraph(values["public_description"]))
            translation.write(dict(values, content_source="manual", source_summary=_("Imported from Excel file: %s") % self.filename))
            if record.state != "draft":
                record.write({"state": "draft"})
            record.message_post(body=_("Imported from Excel file: %s") % self.filename)
            imported |= record
        return {
            "type": "ir.actions.act_window",
            "name": _("Imported SEO Drafts"),
            "res_model": "ab.product.seo",
            "view_mode": "list,form",
            "domain": [("id", "in", imported.ids)],
        }


class AbProductSeoXlsxImportLine(models.TransientModel):
    _name = "ab_product_seo_xlsx_import_line"
    _description = "SEO Excel Import Report Line"
    _order = "id"

    wizard_id = fields.Many2one("ab_product_seo_xlsx_import_wizard", required=True, ondelete="cascade")
    row_number = fields.Char(string="Excel Row", readonly=True)
    product_code = fields.Char(string="Product Code", readonly=True)
    product_name = fields.Char(string="Product Name", readonly=True)
    seo_id = fields.Many2one("ab.product.seo", string="SEO Record", readonly=True)
    status = fields.Selection(
        [("imported", "Imported"), ("not_imported", "Not Imported"), ("error", "Error")],
        string="Status", readonly=True,
    )
    reason = fields.Text(string="Details", readonly=True)
